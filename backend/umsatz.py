"""UMSATZ — jede Buchung mit Betrag, fuer Verwaltung und Steuerberater (25.09.2026, Steve).

Anlass: Michael bucht Rechnungskunden (Actino) von Hand, Stripe bucht online — nirgends
stand ein Euro-Betrag, und eine vertauschte Gutschrift (Verkauf als Kulanz gebucht) fiel
niemandem auf. Jetzt landet jede Buchung als eigene Zeile in der Tabelle `buchungen`:

- Stripe-Paketkauf  -> Webhook checkout.session.completed (Betrag = amount_total)
- Stripe-Abo        -> Webhook invoice.paid (jede bezahlte Abo-Rechnung, Betrag = amount_paid)
- Rechnung          -> Verwaltung: „Credits gutschreiben“ / „Abo zuweisen“, Art „Verkauf auf
                       Rechnung“ mit Betrag (Pflicht) und optionaler Rechnungsnummer
- Bonus (kostenlos) -> Verwaltung, Grund ist Pflicht; zaehlt NICHT zum Umsatz
- Auto-Verlaengerung eines Rechnungs-Abos -> Tageslauf, zum Listenpreis

Betraege stehen als ganze Cent (keine Rundungsfehler), Zeiten in UTC; Tages-, Monats- und
Jahresgrenzen rechnet dieses Modul in deutscher Zeit (Europe/Berlin), damit „heute“ um
Mitternacht in Deutschland wechselt und nicht um 2 Uhr.

Das ist bewusst KEINE Buchhaltung: keine Steuerberechnung, keine Rechnungserstellung. Die
Betraege sind das, was der Kunde zahlt (Stripe: tatsaechlich eingezogen; Rechnung: wie
berechnet). Rechnungen schreibt weiter Actino; die Rechnungsnummer verknuepft beides.
"""

import csv
import io
import logging
import re
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from database import get_db

log = logging.getLogger("umsatz")

ZONE = ZoneInfo("Europe/Berlin")
ARTEN = ("paket", "abo")
WEGE = ("stripe", "rechnung", "bonus")
STATUS = ("ok", "ausstehend", "rueckgelaufen", "storniert")

# Bonus-Gutschriften ueber dieser Menge brauchen eine ausdrueckliche Bestaetigung
# (Steve 25.09.2026: „wie kann man das trotzdem irgendwie eingrenzen?“).
BONUS_GRENZE = 500
# Preisvorschlag fuer freie Mengen auf Rechnung: Satz des kleinsten Pakets (500 = 20 EUR).
CENT_JE_CREDIT_VORSCHLAG = 4
BETRAG_MAX_CENT = 10_000_000  # 100.000 EUR — alles darueber ist ein Tippfehler

# Was zum Umsatz zaehlt: alles ausser Bonus und Ruecklastschrift. Eine SEPA-Lastschrift
# unterwegs zaehlt mit (die Leistung ist erbracht) und wird getrennt ausgewiesen.
_ZAEHLT = "weg != 'bonus' AND status IN ('ok', 'ausstehend')"

WEG_TEXT = {"stripe": "Stripe", "rechnung": "Rechnung", "bonus": "Bonus (kostenlos)"}
ART_TEXT = {"paket": "Credit-Paket", "abo": "Abo"}
STATUS_TEXT = {"ok": "", "ausstehend": "Lastschrift ausstehend",
               "rueckgelaufen": "Rücklastschrift", "storniert": "storniert"}
PLAN_TEXT = {"single": "Single", "team": "Team", "enterprise": "Enterprise"}

_BETRAG = re.compile(r"^\d{1,6}(?:[.,]\d{1,2})?$")


# ─── Umrechnen ───────────────────────────────────────────────────────────

def euro_zu_cent(wert) -> int:
    """'87,50' / '87.50' / 87.5 -> 8750. Nur einfache Schreibweisen (kein Tausenderpunkt),
    damit '1.500' nicht still als 1,50 EUR durchgeht. ValueError bei allem anderen."""
    if isinstance(wert, bool):
        raise ValueError("Betrag fehlt")
    if isinstance(wert, (int, float)):
        cent = round(float(wert) * 100)
    else:
        text = str(wert or "").replace("€", "").replace(" ", "").replace(" ", "")
        if not _BETRAG.match(text):
            raise ValueError("Betrag bitte als Zahl mit höchstens zwei Nachkommastellen, zum Beispiel 87,50")
        cent = round(float(text.replace(",", ".")) * 100)
    if cent < 0 or cent > BETRAG_MAX_CENT:
        raise ValueError("Betrag liegt außerhalb des erlaubten Bereichs")
    return int(cent)


def cent_text(cent) -> str:
    """8750 -> '87,50 €', 123456 -> '1.234,56 €' (deutsche Schreibweise)."""
    s = f"{(int(cent or 0)) / 100:,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".") + " €"


def zahl_text(n) -> str:
    """2500 -> '2.500'."""
    return f"{int(n or 0):,}".replace(",", ".")


def _utc_text(dt: datetime) -> str:
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def lokal(utc_text) -> str:
    """'2026-09-25 06:00:20' (UTC) -> '2026-09-25 08:00' (deutsche Zeit)."""
    if not utc_text:
        return ""
    try:
        dt = datetime.strptime(str(utc_text)[:19], "%Y-%m-%d %H:%M:%S").replace(tzinfo=timezone.utc)
    except ValueError:
        return str(utc_text)[:16]
    return dt.astimezone(ZONE).strftime("%Y-%m-%d %H:%M")


def jetzt_lokal() -> datetime:
    return datetime.now(ZONE)


def zeitraum(jahr: int, monat: int = None, tag: int = None):
    """(von, bis) als UTC-Text fuer ein Jahr, einen Monat oder einen Tag in deutscher Zeit."""
    if tag:
        von = datetime(jahr, monat, tag, tzinfo=ZONE)
        bis = datetime.fromordinal(von.toordinal() + 1).replace(tzinfo=ZONE)
    elif monat:
        von = datetime(jahr, monat, 1, tzinfo=ZONE)
        bis = datetime(jahr + (monat == 12), monat % 12 + 1, 1, tzinfo=ZONE)
    else:
        von = datetime(jahr, 1, 1, tzinfo=ZONE)
        bis = datetime(jahr + 1, 1, 1, tzinfo=ZONE)
    return _utc_text(von), _utc_text(bis)


# ─── Preise ──────────────────────────────────────────────────────────────

def paket_preis_cent(credits: int):
    """Listenpreis eines festen Pakets in Cent, sonst None."""
    import billing
    preis = billing.PAKET_PREISE.get(int(credits or 0))
    return None if preis is None else int(round(preis * 100))


def vorschlag_paket_cent(credits: int) -> int:
    """Vorbelegung im Formular: Listenpreis, bei freien Mengen 4 Cent je Credit."""
    fest = paket_preis_cent(credits)
    return fest if fest is not None else int(credits or 0) * CENT_JE_CREDIT_VORSCHLAG


def abo_preis_cent(plan: str, laufzeit: int) -> int:
    """Listenpreis einer Abo-Laufzeit in Cent (Monatspreis x Monate)."""
    import billing
    if plan not in billing.PLAN_PREISE_EUR or not laufzeit:
        return 0
    return int(round(billing.preis_pro_monat(plan, laufzeit) * int(laufzeit) * 100))


# ─── Buchen ──────────────────────────────────────────────────────────────

def buche(conn, *, konto: dict, art: str, weg: str, credits: int = 0, plan: str = None,
          laufzeit: int = None, betrag_cent: int = 0, rechnungsnummer: str = "", notiz: str = "",
          status: str = "ok", paket_id: int = None, stripe_ref: str = None,
          von: dict = None, von_name: str = "", gebucht_am: str = None):
    """Eine Buchung in der OFFENEN Verbindung anlegen (der Aufrufer committet — so landen
    Paket und Buchung in derselben Transaktion). Rueckgabe: id oder None, wenn die
    stripe_ref schon gebucht war (Stripe wiederholt Webhooks)."""
    if art not in ARTEN or weg not in WEGE or status not in STATUS:
        raise ValueError(f"ungueltige Buchung art={art!r} weg={weg!r} status={status!r}")
    betrag_cent = 0 if weg == "bonus" else int(betrag_cent or 0)
    if betrag_cent < 0 or betrag_cent > BETRAG_MAX_CENT:
        raise ValueError("Betrag ausserhalb des erlaubten Bereichs")
    konto = konto or {}
    felder = {
        "konto_user_id": konto.get("id"),
        "kunde_name": (konto.get("display_name") or "")[:200],
        "kunde_email": (konto.get("email") or "")[:320],
        "art": art, "weg": weg, "credits": int(credits or 0),
        "plan": plan, "laufzeit_monate": int(laufzeit) if laufzeit else None,
        "betrag_cent": betrag_cent,
        "rechnungsnummer": (rechnungsnummer or "").strip()[:100],
        "notiz": (notiz or "").strip()[:500],
        "status": status, "paket_id": paket_id, "stripe_ref": stripe_ref or None,
        "gebucht_von_id": (von or {}).get("id"),
        "gebucht_von_name": ((von or {}).get("display_name") or von_name or "")[:200],
    }
    if gebucht_am:
        felder["gebucht_am"] = gebucht_am
    spalten = ", ".join(felder)
    marken = ", ".join("?" for _ in felder)
    # OR IGNORE nur fuer Stripe (doppelter Webhook = still ueberspringen). Bei Hand-Buchungen
    # soll jeder Fehler laut sein — OR IGNORE wuerde auch Pflichtfeld-Verstoesse verschlucken.
    verb = "INSERT OR IGNORE" if felder["stripe_ref"] else "INSERT"
    cur = conn.execute(f"{verb} INTO buchungen ({spalten}) VALUES ({marken})",
                       list(felder.values()))
    return int(cur.lastrowid) if cur.rowcount else None


def buche_einzeln(**kw):
    """Wie buche(), aber mit eigener Verbindung. Fuer die Stripe-Webhooks: Scheitert die
    Buchung, darf das die Freischaltung NIE verhindern — Fehler werden geloggt."""
    conn = get_db()
    try:
        bid = buche(conn, **kw)
        conn.commit()
        return bid
    except Exception:
        log.exception("Umsatz-Buchung fehlgeschlagen (%s)", {k: kw.get(k) for k in ("art", "weg", "stripe_ref")})
        return None
    finally:
        conn.close()


def setze_status(stripe_ref: str, status: str) -> int:
    """Status einer Stripe-Buchung nachziehen (Lastschrift eingegangen / zurueckgelaufen)."""
    if status not in STATUS or not stripe_ref:
        return 0
    conn = get_db()
    try:
        n = conn.execute("UPDATE buchungen SET status = ? WHERE stripe_ref = ?",
                         (status, stripe_ref)).rowcount
        conn.commit()
        return n
    except Exception:
        log.exception("Umsatz-Status nicht gesetzt (%s -> %s)", stripe_ref, status)
        return 0
    finally:
        conn.close()


def letzte_abo_buchung(konto_id: int):
    conn = get_db()
    try:
        row = conn.execute("SELECT * FROM buchungen WHERE konto_user_id = ? AND art = 'abo' "
                           "ORDER BY gebucht_am DESC, id DESC LIMIT 1", (konto_id,)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def korrigiere(buchung_id: int, *, weg: str, betrag_cent: int, rechnungsnummer: str,
               grund: str, admin: dict, bestaetigt_gross: bool = False) -> dict:
    """Eine von Hand eingetragene Buchung berichtigen (Verkauf <-> Bonus, Betrag, Nummer).

    Stripe-Buchungen sind gesperrt — dort ist Stripe die Wahrheit. Die Credits bleiben
    unveraendert; mitgezogen werden Quelle und Verfall des verknuepften Pakets, damit Report
    und Guthaben dieselbe Sprache sprechen (Verkauf = verfaellt nie, Bonus = 12 Monate ab
    Gutschrift). Jede Korrektur haengt eine Zeile an korrektur_notiz an — nichts wird still
    ueberschrieben."""
    if weg not in ("rechnung", "bonus"):
        raise ValueError("Art muss Verkauf auf Rechnung oder Bonus sein")
    grund = (grund or "").strip()
    if len(grund) < 3:
        raise ValueError("Bitte kurz den Grund der Korrektur angeben")
    betrag_cent = 0 if weg == "bonus" else int(betrag_cent or 0)
    if weg == "rechnung" and betrag_cent <= 0:
        raise ValueError("Ein Verkauf braucht einen Betrag")
    conn = get_db()
    try:
        alt = conn.execute("SELECT * FROM buchungen WHERE id = ?", (buchung_id,)).fetchone()
        if not alt:
            raise LookupError("Buchung nicht gefunden")
        alt = dict(alt)
        if alt["weg"] == "stripe":
            raise ValueError("Stripe-Buchungen lassen sich nicht ändern — dort gilt, was Stripe abgerechnet hat")
        if alt["status"] == "storniert":
            raise ValueError("Eine stornierte Buchung lässt sich nicht mehr berichtigen")
        # Dieselbe Schwelle wie beim Gutschreiben: ein grosser Verkauf wird nicht still zum Geschenk.
        if (weg == "bonus" and alt["weg"] != "bonus" and int(alt["credits"] or 0) > BONUS_GRENZE
                and not bestaetigt_gross):
            raise ValueError(f"Ein Bonus über {BONUS_GRENZE} Credits muss ausdrücklich bestätigt werden")
        stempel = jetzt_lokal().strftime("%d.%m.%Y %H:%M")
        vorher = f"{WEG_TEXT[alt['weg']]}, {cent_text(alt['betrag_cent'])}"
        nachher = f"{WEG_TEXT[weg]}, {cent_text(betrag_cent)}"
        zeile = f"{stempel} {admin.get('display_name') or '?'}: {vorher} → {nachher} ({grund[:200]})"
        protokoll = (alt.get("korrektur_notiz") or "")
        protokoll = (protokoll + "\n" + zeile).strip()[-4000:]
        conn.execute("UPDATE buchungen SET weg = ?, betrag_cent = ?, rechnungsnummer = ?, "
                     "korrigiert_von_name = ?, korrigiert_am = datetime('now'), korrektur_notiz = ? "
                     "WHERE id = ?",
                     (weg, betrag_cent, (rechnungsnummer or "").strip()[:100],
                      (admin.get("display_name") or "")[:200], protokoll, buchung_id))
        if alt.get("paket_id") and alt["art"] == "paket":
            if weg == "rechnung":
                conn.execute("UPDATE quota_pakete SET quelle = 'rechnung', verfaellt_am = NULL "
                             "WHERE id = ?", (alt["paket_id"],))
            else:
                conn.execute("UPDATE quota_pakete SET quelle = 'admin', "
                             "verfaellt_am = COALESCE(verfaellt_am, datetime(erstellt_am, '+12 months')) "
                             "WHERE id = ?", (alt["paket_id"],))
        conn.commit()
        neu = conn.execute("SELECT * FROM buchungen WHERE id = ?", (buchung_id,)).fetchone()
        return darstellen(dict(neu))
    finally:
        conn.close()


def storniere(buchung_id: int, *, grund: str, admin: dict) -> dict:
    """Eine von Hand eingetragene Credit-Gutschrift stornieren (Steve 25.09.2026: Fehlbuchungen
    wie „2.500 statt 250“ muessen rueckgaengig zu machen sein).

    Zurueckgenommen werden nur die NOCH NICHT verbrauchten Credits des Pakets — verbrauchte
    bleiben verbraucht. Die Buchung bleibt sichtbar, steht auf 'storniert' und zaehlt nicht
    mehr zum Umsatz; Grund, Name, Zeit und die zurueckgenommene Menge stehen im Protokoll.
    Auch fuer Stripe (Pruefbericht 25.09.2026): nach einer Erstattung im Stripe-Dashboard wird
    die Buchung hier storniert, sonst bliebe der Umsatz zu hoch. Das Geld selbst erstattet
    NUR Stripe. Bei Abos aendert das Stornieren nur den Umsatz, nicht den Plan (den ueber
    „Abo zuweisen oder aendern“ setzen)."""
    grund = (grund or "").strip()
    if len(grund) < 3:
        raise ValueError("Bitte kurz den Grund für das Stornieren angeben")
    conn = get_db()
    try:
        alt = conn.execute("SELECT * FROM buchungen WHERE id = ?", (buchung_id,)).fetchone()
        if not alt:
            raise LookupError("Buchung nicht gefunden")
        alt = dict(alt)
        if alt["status"] == "storniert":
            raise ValueError("Diese Buchung ist schon storniert")
        zurueck = 0
        if alt["art"] == "paket" and alt.get("paket_id"):
            row = conn.execute("SELECT verbleibend FROM quota_pakete WHERE id = ?", (alt["paket_id"],)).fetchone()
            if row:
                zurueck = int(row["verbleibend"] or 0)
                # Bedingt auf den gelesenen Stand: ein gleichzeitiger Verbrauch darf nicht verloren gehen.
                if not conn.execute("UPDATE quota_pakete SET verbleibend = 0 WHERE id = ? AND verbleibend = ?",
                                    (alt["paket_id"], zurueck)).rowcount:
                    raise ValueError("Das Guthaben hat sich gerade geändert — bitte noch einmal versuchen")
        stempel = jetzt_lokal().strftime("%d.%m.%Y %H:%M")
        if alt["art"] == "paket":
            zeile = (f"{stempel} {admin.get('display_name') or '?'}: storniert, {zahl_text(zurueck)} von "
                     f"{zahl_text(alt['credits'])} Credits zurückgenommen ({grund[:200]})")
        else:
            zeile = f"{stempel} {admin.get('display_name') or '?'}: storniert, Plan unverändert ({grund[:200]})"
        protokoll = ((alt.get("korrektur_notiz") or "") + "\n" + zeile).strip()[-4000:]
        conn.execute("UPDATE buchungen SET status = 'storniert', korrigiert_von_name = ?, "
                     "korrigiert_am = datetime('now'), korrektur_notiz = ? WHERE id = ?",
                     ((admin.get("display_name") or "")[:200], protokoll, buchung_id))
        conn.commit()
        neu = dict(conn.execute("SELECT * FROM buchungen WHERE id = ?", (buchung_id,)).fetchone())
        return {**darstellen(neu), "zurueckgenommen": zurueck}
    finally:
        conn.close()


# ─── Auswerten ───────────────────────────────────────────────────────────

def darstellen(row: dict) -> dict:
    """Buchung fuer Oberflaeche und Export: lokale Zeit, lesbare Texte, keine Interna."""
    return {
        "id": row["id"],
        "gebucht_am": lokal(row["gebucht_am"]),
        "konto_user_id": row.get("konto_user_id"),
        "kunde_name": row.get("kunde_name") or "",
        "kunde_email": row.get("kunde_email") or "",
        "art": row["art"], "art_text": ART_TEXT.get(row["art"], row["art"]),
        "weg": row["weg"], "weg_text": WEG_TEXT.get(row["weg"], row["weg"]),
        "credits": row.get("credits") or 0,
        "plan": row.get("plan"), "plan_text": PLAN_TEXT.get(row.get("plan") or "", ""),
        "laufzeit_monate": row.get("laufzeit_monate"),
        "betrag_cent": row.get("betrag_cent") or 0,
        "betrag_text": cent_text(row.get("betrag_cent") or 0),
        "rechnungsnummer": row.get("rechnungsnummer") or "",
        "notiz": row.get("notiz") or "",
        "status": row.get("status") or "ok",
        "status_text": STATUS_TEXT.get(row.get("status") or "ok", ""),
        "gebucht_von": row.get("gebucht_von_name") or ("Stripe (automatisch)" if row["weg"] == "stripe" else ""),
        "korrigiert": bool(row.get("korrigiert_am")),
        "korrektur_notiz": row.get("korrektur_notiz") or "",
        "paket_rest": row.get("paket_rest"),
    }


def _summe(conn, von: str, bis: str = None) -> int:
    sql = f"SELECT COALESCE(SUM(betrag_cent), 0) FROM buchungen WHERE {_ZAEHLT} AND gebucht_am >= ?"
    werte = [von]
    if bis:
        sql += " AND gebucht_am < ?"
        werte.append(bis)
    return int(conn.execute(sql, werte).fetchone()[0])


def kennzahlen(jetzt: datetime = None) -> dict:
    """Die vier Zahlen oben auf der Umsatz-Seite plus die verschenkten Credits des Monats."""
    jetzt = jetzt or jetzt_lokal()
    tag_von, tag_bis = zeitraum(jetzt.year, jetzt.month, jetzt.day)
    mon_von, mon_bis = zeitraum(jetzt.year, jetzt.month)
    jahr_von, jahr_bis = zeitraum(jetzt.year)
    conn = get_db()
    try:
        bonus = conn.execute(
            "SELECT COUNT(*), COALESCE(SUM(credits), 0) FROM buchungen "
            "WHERE weg = 'bonus' AND gebucht_am >= ? AND gebucht_am < ?", (mon_von, mon_bis)).fetchone()
        ausstehend = conn.execute(
            "SELECT COALESCE(SUM(betrag_cent), 0) FROM buchungen WHERE status = 'ausstehend'").fetchone()[0]
        return {
            "heute_cent": _summe(conn, tag_von, tag_bis),
            "monat_cent": _summe(conn, mon_von, mon_bis),
            "jahr_cent": _summe(conn, jahr_von, jahr_bis),
            "gesamt_cent": _summe(conn, "0000"),
            "monat": f"{jetzt.year:04d}-{jetzt.month:02d}",
            "jahr": jetzt.year,
            "heute": jetzt.strftime("%Y-%m-%d"),
            "bonus_anzahl_monat": int(bonus[0]),
            "bonus_credits_monat": int(bonus[1]),
            "ausstehend_cent": int(ausstehend),
        }
    finally:
        conn.close()


def jahre() -> list:
    """Jahre mit Buchungen (deutsche Zeit) plus das laufende, neuestes zuerst."""
    conn = get_db()
    try:
        erste = conn.execute("SELECT MIN(gebucht_am) FROM buchungen").fetchone()[0]
    finally:
        conn.close()
    aktuell = jetzt_lokal().year
    start = aktuell
    if erste:
        start = min(aktuell, int(lokal(erste)[:4] or aktuell))
    return list(range(aktuell, start - 1, -1))


def monatsuebersicht(jahr: int) -> list:
    """Je Monat des Jahres: Umsatz gesamt, davon Stripe/Rechnung, Anzahl, Bonus-Credits.
    Im laufenden Jahr nur bis zum laufenden Monat (keine leeren Zukunftsmonate)."""
    jetzt = jetzt_lokal()
    bis_monat = jetzt.month if jahr == jetzt.year else (12 if jahr < jetzt.year else 0)
    out = []
    conn = get_db()
    try:
        for m in range(bis_monat, 0, -1):
            von, bis = zeitraum(jahr, m)
            r = conn.execute(
                "SELECT "
                f" COALESCE(SUM(CASE WHEN {_ZAEHLT} THEN betrag_cent END), 0) AS gesamt, "
                f" COALESCE(SUM(CASE WHEN {_ZAEHLT} AND weg = 'stripe' THEN betrag_cent END), 0) AS stripe, "
                f" COALESCE(SUM(CASE WHEN {_ZAEHLT} AND weg = 'rechnung' THEN betrag_cent END), 0) AS rechnung, "
                f" COALESCE(SUM(CASE WHEN {_ZAEHLT} THEN 1 END), 0) AS anzahl, "
                " COALESCE(SUM(CASE WHEN weg = 'bonus' THEN credits END), 0) AS bonus_credits "
                "FROM buchungen WHERE gebucht_am >= ? AND gebucht_am < ?", (von, bis)).fetchone()
            out.append({"monat": f"{jahr:04d}-{m:02d}", "gesamt_cent": int(r["gesamt"]),
                        "stripe_cent": int(r["stripe"]), "rechnung_cent": int(r["rechnung"]),
                        "anzahl": int(r["anzahl"]), "bonus_credits": int(r["bonus_credits"])})
    finally:
        conn.close()
    return out


def liste(jahr: int = None, monat: int = None, auswahl: str = "alle", konto_id: int = None) -> list:
    """Buchungen eines Jahres/Monats (oder eines Kontos), neueste zuerst.
    auswahl: 'alle' | 'verkauf' (Stripe + Rechnung) | 'bonus'."""
    bedingungen, werte = [], []
    if jahr:
        von, bis = zeitraum(int(jahr), int(monat) if monat else None)
        bedingungen.append("b.gebucht_am >= ? AND b.gebucht_am < ?")
        werte += [von, bis]
    if konto_id is not None:
        bedingungen.append("b.konto_user_id = ?")
        werte.append(int(konto_id))
    if auswahl == "verkauf":
        bedingungen.append("b.weg != 'bonus'")
    elif auswahl == "bonus":
        bedingungen.append("b.weg = 'bonus'")
    sql = ("SELECT b.*, p.verbleibend AS paket_rest FROM buchungen b "
           "LEFT JOIN quota_pakete p ON p.id = b.paket_id")
    if bedingungen:
        sql += " WHERE " + " AND ".join(bedingungen)
    sql += " ORDER BY b.gebucht_am DESC, b.id DESC"
    conn = get_db()
    try:
        return [darstellen(dict(r)) for r in conn.execute(sql, werte).fetchall()]
    finally:
        conn.close()


# ─── Export ──────────────────────────────────────────────────────────────

SPALTEN = ["Datum", "Kunde", "E-Mail", "Art", "Weg", "Credits", "Plan", "Laufzeit (Monate)",
           "Betrag (EUR)", "Rechnungsnummer", "Status", "Notiz", "Eingetragen von", "Korrekturen"]


_STEUERZEICHEN = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _sicher(wert):
    """Formel-Einschleusung verhindern (CSV/Excel-Injection): Ein Kundenname wie
    '=HYPERLINK(...)' darf beim Oeffnen im Tabellenprogramm nie als Formel laufen.
    Steuerzeichen fliegen raus — openpyxl bricht sonst den ganzen Export ab."""
    if isinstance(wert, str):
        wert = _STEUERZEICHEN.sub("", wert)
        if wert[:1] in ("=", "+", "-", "@", "\t", "\r"):
            return "'" + wert
    return wert


def _zeile(b: dict) -> list:
    return [b["gebucht_am"], b["kunde_name"], b["kunde_email"], b["art_text"], b["weg_text"],
            b["credits"] or "", b["plan_text"], b["laufzeit_monate"] or "",
            b["betrag_cent"] / 100, b["rechnungsnummer"], b["status_text"], b["notiz"],
            b["gebucht_von"], b["korrektur_notiz"].replace("\n", " | ")]


def export_csv(buchungen: list) -> bytes:
    """Semikolon und Dezimalkomma — so oeffnet ein deutsches Excel die Datei ohne Import-
    Dialog; das BOM sorgt fuer richtige Umlaute."""
    puffer = io.StringIO()
    w = csv.writer(puffer, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    w.writerow(SPALTEN)
    for b in buchungen:
        z = [_sicher(x) for x in _zeile(b)]
        z[8] = f"{b['betrag_cent'] / 100:.2f}".replace(".", ",")
        w.writerow(z)
    return ("﻿" + puffer.getvalue()).encode("utf-8")


def export_xlsx(buchungen: list, titel: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Buchungen"
    ws.append(SPALTEN)
    for zelle in ws[1]:
        zelle.font = Font(bold=True)
    for b in buchungen:
        ws.append([_sicher(x) for x in _zeile(b)])
    letzte = ws.max_row
    for r in range(2, letzte + 1):
        ws.cell(row=r, column=9).number_format = '#,##0.00 "€"'
    # Summenzeile: nur was zum Umsatz zaehlt (kein Bonus, keine Ruecklastschrift) — als
    # fester Wert, damit die Datei auch ohne Formel-Neuberechnung dieselbe Zahl zeigt.
    summe = sum(b["betrag_cent"] for b in buchungen
                if b["weg"] != "bonus" and b["status"] in ("ok", "ausstehend")) / 100
    ws.append([])
    ws.append(["Umsatz " + titel, "", "", "", "", "", "", "", summe])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=9).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=9).number_format = '#,##0.00 "€"'
    ws.freeze_panes = "A2"
    if letzte > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(SPALTEN))}{letzte}"
    breiten = [17, 26, 30, 13, 18, 9, 11, 10, 13, 18, 20, 40, 22, 40]
    for i, b in enumerate(breiten, start=1):
        ws.column_dimensions[get_column_letter(i)].width = b
    puffer = io.BytesIO()
    wb.save(puffer)
    return puffer.getvalue()
