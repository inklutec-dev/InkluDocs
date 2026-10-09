#!/usr/bin/env python3
"""Fiktives Programmheft „Gartenfest 2026“ (Nachbarschaftsverein Musterstadt e. V.) als DOCX mit SIEBEN Bildern fuer die
Fall-Matrix „Alt-Texte: Feld = Datei“ (09.10.2026). Alle Namen und Zahlen sind erfunden.
B1..B4 tragen einen Alt-Text (docPr descr) und eine Bildunterschrift, B5 ist vom Autor als dekorativ markiert
(adec:decorative), B6 hat keinen Alt-Text (aber eine Bildunterschrift),
B7 (Vereinslogo) traegt einen Alt-Text ohne Unterschrift.
Aufruf: python make_gartenfest.py <zielordner>; die DOCX wandelt der InkluDocs-Umwandler (LibreOffice, PDF/UA)
in tests/fixtures/gartenfest_feld_datei.pdf um (pdfua_export.konvertiere, 09.10.2026)."""
import math
import os
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw

ZIEL = sys.argv[1]
os.makedirs(ZIEL, exist_ok=True)
DOCX = os.path.join(ZIEL, "Gartenfest-2026_Programm.docx")
W, H = 900, 560


def bild(name, malen, groesse=(W, H)):
    im = Image.new("RGB", groesse, "white")
    malen(im, ImageDraw.Draw(im))
    p = os.path.join(ZIEL, name)
    im.save(p)
    return p


def b1(im, d):   # Sonnenblumen: Himmel-Verlauf, gelbe Kreise unten
    for y in range(H):
        d.line((0, y, W, y), fill=(40, 90 + int(120 * y / H), 220))
    for i, x in enumerate(range(80, W, 160)):
        d.ellipse((x - 55, 330 - 30 * (i % 2), x + 55, 440 - 30 * (i % 2)), fill=(250, 200, 20), outline=(120, 70, 10), width=8)
        d.rectangle((x - 6, 440, x + 6, H), fill=(40, 120, 40))


def b2(im, d):   # Kuchenbuffet: braune Bloecke im Raster
    for r in range(3):
        for c in range(4):
            x, y = 40 + c * 215, 40 + r * 170
            d.rectangle((x, y, x + 190, y + 140), fill=(140 + 30 * ((r + c) % 2), 80, 30))


def b3(im, d):   # Wimpel: diagonale gruene Streifen
    for k in range(-H, W, 70):
        d.polygon([(k, 0), (k + 35, 0), (k + 35 + H, H), (k + H, H)], fill=(30, 140, 60))


def b4(im, d):   # Plakat: grosses rot-weisses Schachbrett, links oben schwarz
    for r in range(4):
        for c in range(6):
            if (r + c) % 2 == 0:
                d.rectangle((c * 150, r * 140, c * 150 + 150, r * 140 + 140), fill=(200, 20, 20))
    d.rectangle((0, 0, 300, 280), fill=(0, 0, 0))


def b5(im, d):   # Zierlinie (dekorativ)
    w, h = im.size
    for x in range(0, w, 2):
        y = int(h / 2 + (h / 3) * math.sin(x / 25.0))
        d.ellipse((x - 3, y - 3, x + 3, y + 3), fill=(120, 60, 140))


def b6(im, d):   # Lageplan: konzentrische Ringe, rechts hell
    for i, rad in enumerate(range(520, 20, -50)):
        g = 30 + 20 * i
        d.ellipse((W - 200 - rad, H / 2 - rad, W - 200 + rad, H / 2 + rad), fill=(g, g, 255 - g))


def b7(im, d):   # Vereinslogo: gruenes Blatt auf hellem Kreis
    w, h = im.size
    d.ellipse((10, 10, w - 10, h - 10), fill=(235, 245, 225))
    d.polygon([(w * 0.2, h * 0.8), (w * 0.5, h * 0.15), (w * 0.8, h * 0.8)], fill=(20, 110, 50))
    d.line((w * 0.5, h * 0.25, w * 0.5, h * 0.9), fill=(90, 60, 20), width=12)


BILDER = [bild("b1_sonnenblumen.png", b1), bild("b2_kuchenbuffet.png", b2), bild("b3_wimpel.png", b3),
          bild("b4_plakat.png", b4), bild("b5_zierlinie.png", b5, (1200, 90)), bild("b6_lageplan.png", b6), bild("b7_logo.png", b7, (420, 420))]

doc = Document()
rpr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
lang = rpr.find(qn("w:lang"))
if lang is None:
    lang = OxmlElement("w:lang")
    rpr.append(lang)
lang.set(qn("w:val"), "de-DE")
st = doc.styles["Normal"]
st.font.name = "Liberation Sans"
st.font.size = Pt(11)
for s in ("Heading 1", "Heading 2"):
    doc.styles[s].font.name = "Liberation Sans"
    doc.styles[s].font.color.rgb = RGBColor(0x1B, 0x2A, 0x4A)
cp = doc.core_properties
cp.title = "Gartenfest 2026 – Programm"
cp.author = "Nachbarschaftsverein Musterstadt e. V. (fiktiv)"
cp.language = "de-DE"
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)


def P(text):
    return doc.add_paragraph(text)


def bild_einfuegen(pfad, breite, alt=None, titel=None, dekorativ=False, unterschrift=None):
    pb = doc.add_paragraph()
    pb.add_run().add_picture(pfad, width=breite)
    for el in pb._p.iter(qn("wp:docPr")):
        if alt is not None:
            el.set("descr", alt)
        if titel:
            el.set("title", titel)
        if dekorativ:
            ext = parse_xml(
                "<a:extLst xmlns:a=\"http://schemas.openxmlformats.org/drawingml/2006/main\">"
                "<a:ext uri=\"{C183D7F6-B498-43B3-948B-1728B52AA6E4}\">"
                "<adec:decorative xmlns:adec=\"http://schemas.microsoft.com/office/drawing/2017/decorative\" val=\"1\"/>"
                "</a:ext></a:extLst>")
            el.append(ext)
    if unterschrift:
        doc.add_paragraph(unterschrift, style="Caption")


doc.add_heading("Gartenfest 2026 – Programm", level=1)
P("Nachbarschaftsverein Musterstadt e. V. – erfundenes Testdokument, alle Namen und Zahlen sind ausgedacht.")
doc.add_heading("Willkommen im Vereinsgarten", level=2)
P("Am Samstag, 12. September 2026, feiern wir ab 14 Uhr unser Gartenfest hinter dem Vereinsheim. "
  "Es gibt Musik, Spiele für Kinder und ein großes Kuchenbuffet.")
bild_einfuegen(BILDER[0], Cm(11), alt="Gelbe Sonnenblumen vor blauem Himmel am Gartenzaun des Vereinsgartens.",
               unterschrift="Abbildung 1: Sonnenblumen im Vereinsgarten, Sommer 2026.")
P("Die Sonnenblumen hat die Kindergruppe im Mai gesät. Sie sind inzwischen über zwei Meter hoch.")
bild_einfuegen(BILDER[1], Cm(11), alt="Kuchenbuffet mit zwölf Blechkuchen auf einem langen Tisch.",
               unterschrift="Abbildung 2: Das Kuchenbuffet vom letzten Jahr.")
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
doc.add_heading("Spiele und Schmuck", level=2)
P("Für die Kinder gibt es Sackhüpfen, Dosenwerfen und eine Schatzsuche durch den Garten.")
bild_einfuegen(BILDER[2], Cm(11), alt="Grüne Wimpelketten über der Festwiese.",
               unterschrift="Abbildung 3: Wimpel über der Festwiese.")
bild_einfuegen(BILDER[4], Cm(15), dekorativ=True)
P("Die Wimpel nähen die Seniorinnen und Senioren der Handarbeitsgruppe.")
bild_einfuegen(BILDER[3], Cm(11), alt="Plakat zum Gartenfest mit der Telefonnummer der Vorsitzenden.",
               unterschrift="Abbildung 4: Plakat zum Gartenfest.")
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
doc.add_heading("Anfahrt", level=2)
P("Der Vereinsgarten liegt hinter dem Vereinsheim am Lindenplatz 3 in Musterstadt. Fahrräder können im Hof stehen.")
bild_einfuegen(BILDER[5], Cm(11), unterschrift="Abbildung 5: Lageplan des Vereinsgartens.")
P("Bei Fragen hilft der Vorstand gern weiter. Wir freuen uns auf euch!")
bild_einfuegen(BILDER[6], Cm(3), alt="Logo des Nachbarschaftsvereins Musterstadt.")
doc.save(DOCX)
print(DOCX)
