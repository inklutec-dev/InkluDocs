"""Umsatz und Verwaltung (Steve 25.09.2026) — auf einer WEGWERF-Datenbank, nie auf echten Daten:
    docker exec inkludocs-staging python3 -m unittest /app/tests/test_umsatz.py -v

Prueft: Betraege und Zeitgrenzen, Paket + Buchung in einer Transaktion, doppelte Stripe-
Webhooks, Summen (Bonus und Ruecklastschrift zaehlen nicht), Berichtigen, Kontoloeschung
(Buchungen bleiben), Export (Formel-Einschleusung, Summenzeile) und die Endpunkte samt
Rechten (Nur-Einsicht darf lesen, aber nicht buchen).
"""
import csv
import io
import os
import sys
import tempfile
import unittest
from datetime import datetime

_TMP = tempfile.mkdtemp(prefix="umsatz-test-")
os.environ["INKLUDOCS_DB"] = os.path.join(_TMP, "test.db")

HERE = os.path.dirname(os.path.abspath(__file__))
for kandidat in ("/app", os.path.join(os.path.dirname(HERE), "backend")):
    if os.path.isdir(kandidat) and kandidat not in sys.path:
        sys.path.insert(0, kandidat)

import database  # noqa: E402

assert database.DB_PATH.startswith(_TMP), "Test darf nur auf der Wegwerf-Datenbank laufen"
database.init_db()

import billing  # noqa: E402
import umsatz  # noqa: E402
from umsatz import ZONE  # noqa: E402


def _konto(email, name, admin=0, stufe="full"):
    uid = database.create_user(email, "geheim-12345", name)
    if admin:
        database.set_user_admin(uid, 1, stufe)
    return database.get_user_by_id(uid)


class TestUmrechnen(unittest.TestCase):
    def test_euro_zu_cent(self):
        self.assertEqual(umsatz.euro_zu_cent("87,50"), 8750)
        self.assertEqual(umsatz.euro_zu_cent("87.5"), 8750)
        self.assertEqual(umsatz.euro_zu_cent(" 20 € "), 2000)
        self.assertEqual(umsatz.euro_zu_cent(150), 15000)
        for falsch in ("1.500,00", "abc", "", "-5", "12,345", None, True):
            with self.assertRaises(ValueError, msg=repr(falsch)):
                umsatz.euro_zu_cent(falsch)
        with self.assertRaises(ValueError):
            umsatz.euro_zu_cent(100001)

    def test_texte(self):
        self.assertEqual(umsatz.cent_text(8750), "87,50 €")
        self.assertEqual(umsatz.cent_text(123456), "1.234,56 €")
        self.assertEqual(umsatz.zahl_text(2500), "2.500")

    def test_zeitgrenzen_deutsche_zeit(self):
        # September 2026 beginnt in Deutschland (Sommerzeit) um 22:00 UTC am 31. August.
        self.assertEqual(umsatz.zeitraum(2026, 9), ("2026-08-31 22:00:00", "2026-09-30 22:00:00"))
        # Tag mit Zeitumstellung (25.10.2026: 25 Stunden).
        self.assertEqual(umsatz.zeitraum(2026, 10, 25), ("2026-10-24 22:00:00", "2026-10-25 23:00:00"))
        self.assertEqual(umsatz.zeitraum(2026), ("2025-12-31 23:00:00", "2026-12-31 23:00:00"))
        self.assertEqual(umsatz.lokal("2026-09-25 06:00:20"), "2026-09-25 08:00")

    def test_preise(self):
        self.assertEqual(umsatz.vorschlag_paket_cent(2500), 8750)
        self.assertEqual(umsatz.vorschlag_paket_cent(150), 600)
        self.assertEqual(umsatz.abo_preis_cent("single", 6), 5970)
        self.assertEqual(umsatz.abo_preis_cent("free", 6), 0)


class TestBuchen(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.kunde = _konto("kunde.umsatz@example.invalid", "Testkunde Umsatz")
        cls.admin = _konto("admin.umsatz@example.invalid", "Testadmin Umsatz", admin=1)

    def _buchungen(self):
        return umsatz.liste(konto_id=self.kunde["id"])

    def test_1_paket_und_buchung_in_einer_transaktion(self):
        pid = billing.schenke_credits(self.kunde["id"], 2500, notiz="Test", quelle="rechnung",
                                      verfall_monate=None,
                                      buchung={"konto": self.kunde, "weg": "rechnung", "betrag_cent": 8750,
                                               "rechnungsnummer": "R-1", "notiz": "Test", "von": self.admin})
        b = [x for x in self._buchungen() if x["credits"] == 2500][0]
        self.assertEqual((b["weg"], b["betrag_cent"], b["rechnungsnummer"], b["gebucht_von"]),
                         ("rechnung", 8750, "R-1", "Testadmin Umsatz"))
        conn = database.get_db()
        p = dict(conn.execute("SELECT * FROM quota_pakete WHERE id = ?", (pid,)).fetchone())
        conn.close()
        self.assertEqual((p["quelle"], p["verfaellt_am"]), ("rechnung", None))

    def test_2_fehlerhafte_buchung_legt_kein_paket_an(self):
        conn = database.get_db()
        vorher = conn.execute("SELECT COUNT(*) FROM quota_pakete").fetchone()[0]
        conn.close()
        with self.assertRaises(ValueError):
            billing.schenke_credits(self.kunde["id"], 10, quelle="rechnung", verfall_monate=None,
                                    buchung={"konto": self.kunde, "weg": "falsch", "betrag_cent": 1})
        conn = database.get_db()
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM quota_pakete").fetchone()[0], vorher)
        conn.close()

    def test_3_stripe_doppelt_bucht_einmal(self):
        erst = umsatz.buche_einzeln(konto=self.kunde, art="paket", weg="stripe", credits=500,
                                    betrag_cent=2000, stripe_ref="cs_test_doppelt", status="ausstehend")
        zweit = umsatz.buche_einzeln(konto=self.kunde, art="paket", weg="stripe", credits=500,
                                     betrag_cent=2000, stripe_ref="cs_test_doppelt")
        self.assertIsNotNone(erst)
        self.assertIsNone(zweit)
        self.assertEqual(umsatz.setze_status("cs_test_doppelt", "ok"), 1)

    def test_4_summen(self):
        umsatz.buche_einzeln(konto=self.kunde, art="paket", weg="bonus", credits=250, betrag_cent=999,
                             notiz="Geschenk")
        umsatz.buche_einzeln(konto=self.kunde, art="paket", weg="stripe", credits=500, betrag_cent=2000,
                             stripe_ref="cs_test_rueck")
        umsatz.setze_status("cs_test_rueck", "rueckgelaufen")
        k = umsatz.kennzahlen()
        # 87,50 (Rechnung) + 20,00 (Stripe) — Bonus (auch mit falschem Betrag) und Ruecklastschrift nicht.
        self.assertEqual(k["monat_cent"], 10750)
        self.assertEqual(k["gesamt_cent"], 10750)
        self.assertEqual(k["heute_cent"], 10750)
        self.assertGreaterEqual(k["bonus_credits_monat"], 250)
        jetzt = datetime.now(ZONE)
        monat = [m for m in umsatz.monatsuebersicht(jetzt.year) if m["monat"] == f"{jetzt.year}-{jetzt.month:02d}"][0]
        self.assertEqual((monat["gesamt_cent"], monat["stripe_cent"], monat["rechnung_cent"]), (10750, 2000, 8750))
        self.assertTrue(all(b["betrag_cent"] == 0 for b in umsatz.liste(jetzt.year, jetzt.month, "bonus")))
        self.assertTrue(all(b["weg"] != "bonus" for b in umsatz.liste(jetzt.year, jetzt.month, "verkauf")))

    def test_5_berichtigen(self):
        pid = billing.schenke_credits(self.kunde["id"], 2500, notiz="vertauscht", quelle="admin",
                                      verfall_monate=12,
                                      buchung={"konto": self.kunde, "weg": "bonus", "notiz": "vertauscht"})
        b = [x for x in self._buchungen() if x["notiz"] == "vertauscht"][0]
        neu = umsatz.korrigiere(b["id"], weg="rechnung", betrag_cent=8750, rechnungsnummer="R-2",
                                grund="war ein Verkauf", admin={"display_name": "Testadmin Umsatz"})
        self.assertEqual((neu["weg"], neu["betrag_cent"], neu["korrigiert"]), ("rechnung", 8750, True))
        self.assertIn("Bonus (kostenlos), 0,00 € → Rechnung, 87,50 €", neu["korrektur_notiz"])
        conn = database.get_db()
        p = dict(conn.execute("SELECT quelle, verfaellt_am FROM quota_pakete WHERE id = ?", (pid,)).fetchone())
        conn.close()
        self.assertEqual((p["quelle"], p["verfaellt_am"]), ("rechnung", None))
        with self.assertRaises(ValueError):
            umsatz.korrigiere(b["id"], weg="bonus", betrag_cent=0, rechnungsnummer="", grund="x",
                              admin={"display_name": "A"})
        stripe_b = [x for x in self._buchungen() if x["weg"] == "stripe"][0]
        with self.assertRaises(ValueError):
            umsatz.korrigiere(stripe_b["id"], weg="bonus", betrag_cent=0, rechnungsnummer="",
                              grund="geht nicht", admin={"display_name": "A"})

    def test_6_export(self):
        feind = _konto("feind@example.invalid", "=HYPERLINK(\"http://x\")")
        umsatz.buche_einzeln(konto=feind, art="paket", weg="rechnung", credits=500, betrag_cent=2000,
                             notiz="+SUMME(A1)")
        jetzt = datetime.now(ZONE)
        zeilen = umsatz.liste(jetzt.year, jetzt.month)
        text = umsatz.export_csv(zeilen).decode("utf-8-sig")
        daten = list(csv.reader(io.StringIO(text), delimiter=";"))
        self.assertEqual(daten[0][:3], ["Datum", "Kunde", "E-Mail"])
        feind_zeile = [z for z in daten if "feind@example.invalid" in z][0]
        self.assertTrue(feind_zeile[1].startswith("'="))
        self.assertTrue(feind_zeile[11].startswith("'+"))
        self.assertEqual(feind_zeile[8], "20,00")
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(umsatz.export_xlsx(zeilen, "Test")))
        ws = wb.active
        letzte = [c.value for c in ws[ws.max_row]]
        erwartet = sum(b["betrag_cent"] for b in zeilen if b["weg"] != "bonus"
                       and b["status"] in ("ok", "ausstehend")) / 100
        self.assertEqual(letzte[0], "Umsatz Test")
        self.assertAlmostEqual(letzte[8], erwartet)
        self.assertTrue(all(not (isinstance(c.value, str) and c.value.startswith("=")) for row in ws.iter_rows() for c in row))

    def test_7_loeschen_behaelt_buchungen(self):
        weg = _konto("weg@example.invalid", "Wird Geloescht")
        billing.schenke_credits(weg["id"], 500, quelle="rechnung", verfall_monate=None,
                                buchung={"konto": weg, "weg": "rechnung", "betrag_cent": 2000})
        database.delete_user_data(weg["id"])
        conn = database.get_db()
        r = dict(conn.execute("SELECT * FROM buchungen WHERE kunde_email = 'weg@example.invalid'").fetchone())
        conn.close()
        self.assertIsNone(r["konto_user_id"])
        self.assertIsNone(r["paket_id"])
        self.assertEqual((r["kunde_name"], r["betrag_cent"]), ("Wird Geloescht", 2000))


class TestEndpunkte(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            import main
            from fastapi.testclient import TestClient
        except Exception as e:  # ausserhalb des Containers (kein /app/frontend)
            raise unittest.SkipTest(f"main nicht ladbar: {e}")
        cls.main = main
        cls.voll = _konto("voll.ep@example.invalid", "Voll Admin", admin=1)
        cls.sicht = _konto("sicht.ep@example.invalid", "Sicht Admin", admin=1, stufe="view")
        cls.kunde = _konto("kunde.ep@example.invalid", "Kunde Endpunkt")
        cls.c_voll = TestClient(main.app)
        cls.c_voll.cookies.set("token", main.create_token(cls.voll["id"], cls.voll["email"], 1))
        cls.c_sicht = TestClient(main.app)
        cls.c_sicht.cookies.set("token", main.create_token(cls.sicht["id"], cls.sicht["email"], 1))
        cls.c_kunde = TestClient(main.app)
        cls.c_kunde.cookies.set("token", main.create_token(cls.kunde["id"], cls.kunde["email"], 0))
        cls.c_gast = TestClient(main.app)

    def test_rechte(self):
        self.assertEqual(self.c_gast.get("/api/admin/umsatz").status_code, 401)
        self.assertEqual(self.c_kunde.get("/api/admin/umsatz").status_code, 403)
        self.assertEqual(self.c_kunde.get("/api/admin/kunden").status_code, 403)
        self.assertEqual(self.c_kunde.get("/api/admin/umsatz/export").status_code, 403)
        self.assertEqual(self.c_sicht.get("/api/admin/umsatz").status_code, 200)
        r = self.c_sicht.post(f"/api/admin/users/{self.kunde['id']}/pakete",
                              json={"groesse": 500, "art": "bonus", "notiz": "nein"})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(self.c_sicht.post("/api/admin/buchungen/1/korrektur",
                                           json={"art": "bonus", "grund": "nein"}).status_code, 403)

    def test_gutschrift_pflichtangaben(self):
        url = f"/api/admin/users/{self.kunde['id']}/pakete"
        self.assertEqual(self.c_voll.post(url, json={"groesse": 500}).status_code, 400)          # Art fehlt
        self.assertEqual(self.c_voll.post(url, json={"groesse": 500, "art": "verkauf"}).status_code, 400)
        self.assertEqual(self.c_voll.post(url, json={"groesse": 500, "art": "verkauf", "betrag": "0"}).status_code, 400)
        self.assertEqual(self.c_voll.post(url, json={"groesse": 500, "art": "bonus"}).status_code, 400)
        self.assertEqual(self.c_voll.post(url, json={"groesse": 600, "art": "bonus", "notiz": "zu gross"}).status_code, 400)
        r = self.c_voll.post(url, json={"groesse": 600, "art": "bonus", "notiz": "Dankeschön", "bestaetigt_gross": True})
        self.assertEqual(r.status_code, 200, r.text)
        r = self.c_voll.post(url, json={"groesse": 2500, "art": "verkauf", "betrag": "87,50", "rechnungsnummer": "A-7"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("2.500 Credits", r.json()["message"])
        b = self.c_voll.get(f"/api/admin/kunden/{self.kunde['id']}/buchungen").json()["buchungen"]
        verkauf = [x for x in b if x["rechnungsnummer"] == "A-7"][0]
        self.assertEqual((verkauf["weg"], verkauf["betrag_cent"], verkauf["gebucht_von"]), ("rechnung", 8750, "Voll Admin"))

    def test_abo_mit_betrag(self):
        url = f"/api/admin/users/{self.kunde['id']}/plan"
        self.assertEqual(self.c_voll.post(url, json={"plan": "single", "laufzeit_monate": 6}).status_code, 400)
        r = self.c_voll.post(url, json={"plan": "single", "laufzeit_monate": 6, "art": "verkauf",
                                        "betrag": "59,70", "auto_verlaengerung": True})
        self.assertEqual(r.status_code, 200, r.text)
        b = self.c_voll.get(f"/api/admin/kunden/{self.kunde['id']}/buchungen").json()["buchungen"]
        abo = [x for x in b if x["art"] == "abo"][0]
        self.assertEqual((abo["plan"], abo["laufzeit_monate"], abo["betrag_cent"]), ("single", 6, 5970))
        # Free (Abo beenden) bucht nichts und braucht keine Art.
        self.assertEqual(self.c_voll.post(url, json={"plan": "free"}).status_code, 200)

    def test_kundenliste(self):
        d = self.c_voll.get("/api/admin/kunden", params={"q": "kunde.ep"}).json()
        self.assertEqual(d["gesamt"], 1)
        self.assertEqual(d["kunden"][0]["email"], "kunde.ep@example.invalid")
        self.assertNotIn("password_hash", d["kunden"][0])
        self.assertEqual(self.c_voll.get("/api/admin/kunden", params={"filter": "boese"}).status_code, 400)
        d = self.c_voll.get("/api/admin/kunden", params={"filter": "admins"}).json()
        self.assertTrue(all(k["admin"] for k in d["kunden"]))
        d = self.c_voll.get("/api/admin/kunden", params={"seite": 999}).json()
        self.assertEqual(d["seite"], d["seiten"])

    def test_seiten_der_kundenliste(self):
        for i in range(30):
            _konto(f"seite{i:02d}@blaettern.invalid", f"Blaettern {i:02d}")
        d = self.c_voll.get("/api/admin/kunden", params={"q": "blaettern.invalid"}).json()
        self.assertEqual((d["gesamt"], d["seiten"], len(d["kunden"])), (30, 2, 25))
        d2 = self.c_voll.get("/api/admin/kunden", params={"q": "blaettern.invalid", "seite": 2}).json()
        self.assertEqual(len(d2["kunden"]), 5)
        self.assertFalse({k["id"] for k in d["kunden"]} & {k["id"] for k in d2["kunden"]})

    def test_umsatz_und_export(self):
        d = self.c_voll.get("/api/admin/umsatz").json()
        jetzt = datetime.now(ZONE)
        self.assertEqual((d["jahr"], d["monat"]), (jetzt.year, jetzt.month))
        self.assertEqual(self.c_voll.get("/api/admin/umsatz", params={"jahr": 1999}).status_code, 400)
        self.assertEqual(self.c_voll.get("/api/admin/umsatz", params={"jahr": jetzt.year, "monat": 13}).status_code, 400)
        r = self.c_voll.get("/api/admin/umsatz/export", params={"format": "xlsx", "jahr": jetzt.year})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"PK"))
        self.assertIn(f"InkluDocs-Umsatz-{jetzt.year}.xlsx", r.headers["content-disposition"])
        r = self.c_voll.get("/api/admin/umsatz/export", params={"format": "csv", "jahr": jetzt.year, "monat": jetzt.month})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.c_voll.get("/api/admin/umsatz/export", params={"format": "exe"}).status_code, 400)

    def test_api_konten_und_seiten(self):
        self.assertEqual(self.c_voll.get("/api/admin/api-konten").status_code, 200)
        for pfad in ("/verwaltung/kunden", "/verwaltung/umsatz", "/verwaltung/api",
                     "/verwaltung/einstellungen", f"/verwaltung/kunden/{self.kunde['id']}"):
            r = self.c_voll.get(pfad)
            self.assertEqual(r.status_code, 200, pfad)
            self.assertIn('aria-current="page"', r.text, pfad)
        r = self.c_voll.get("/benutzer", follow_redirects=False)
        self.assertEqual((r.status_code, r.headers["location"]), (301, "/verwaltung/kunden"))
        r = self.c_voll.get("/benutzer/report/7", follow_redirects=False)
        self.assertEqual(r.headers["location"], "/verwaltung/kunden/7")


if __name__ == "__main__":
    unittest.main()
