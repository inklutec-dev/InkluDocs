#!/usr/bin/env python3
"""Rauchtest Staging, Livestream zum Anklemmen (17.10.2026) mit dem ECHTEN Modell: Steves Word-Kette und der Gegenfall.

Fall A „ganze Kette“: fiktives Word-Dokument mit zwei Bildern, Agentenansicht, „Alt-Texte, dann barrierefreie PDF, als einen Auftrag“
→ Auftragskarte → Knopf. Das Arbeitsfenster muss Station für Station mitgehen: Alt-Texte (Karte je Bild), Umwandeln (Ansicht
Dokument), am Ende „Meine Ablage“ beim fertigen PDF.
Fall B „nur Alt-Texte und exportieren“: neues Word-Dokument, „nur Alt-Texte und die Word-Datei“ → nur diese Stationen, NIE Prüfung,
Tagging, Übersetzung oder Umwandeln.
Ein Rekorder schreibt alle 300 ms mit, was im Fenster steht (Seite, Ansicht) und was das Live-Fenster meldet (Art/Station).
Fotos /tmp/kette_*.png. Kostet einige Credits des Staging-Kontos und wenige Cent KI. Keine Mail (Haken aus). Projekte und die neuen
Ablage-Einträge danach gelöscht, Ansicht des Kontos wie vorher. Aufruf (vor jedem Einspielen auf Staging, Playwright): python3 tests/e2e/live_kette_rauch.py <mail> <passwort> [A|B|AB]"""
import io
import os
import sys
import time

from docx import Document
from docx.shared import Cm
from PIL import Image, ImageDraw
from playwright.sync_api import sync_playwright

B = os.environ.get("INKLUDOCS_BASIS", "https://staging.inkludocs.inklutec.de")   # nie gegen Production
KARTE = ".inkluagent-formular"
FAELLE = sys.argv[3] if len(sys.argv) > 3 else "AB"
ok = fehler = 0


def check(n, c, info=""):
    global ok, fehler
    ok, fehler = (ok + 1, fehler) if c else (ok, fehler + 1)
    print(("  OK   " if c else "  FEHLT ") + n + ("" if c else " -- " + str(info)[:700]))


def bild(form):
    b = Image.new("RGB", (480, 300), (255, 255, 255))
    d = ImageDraw.Draw(b)
    if form == "haus":
        d.rectangle((150, 160, 330, 250), fill=(240, 220, 170))
        d.polygon([(240, 80), (120, 160), (360, 160)], fill=(160, 40, 40))
    else:
        d.rectangle((225, 180, 255, 260), fill=(110, 70, 40))
        d.ellipse((150, 60, 330, 200), fill=(60, 150, 70))
    p = io.BytesIO()
    b.save(p, "PNG")
    p.seek(0)
    return p


def word(titel):
    doc = Document()
    doc.add_heading(titel, 1)
    doc.add_paragraph("Unser neues Vereinsheim ist fertig. Alle Mitglieder sind herzlich eingeladen (fiktiv).")
    doc.add_picture(bild("haus"), width=Cm(8))
    doc.add_paragraph("Im Garten steht jetzt ein Apfelbaum, gepflanzt beim Sommerfest (fiktiv).")
    doc.add_picture(bild("baum"), width=Cm(8))
    p = io.BytesIO()
    doc.save(p)
    return p.getvalue()


REKORDER = """() => { window._aufnahme = []; window._ziele = [];
    clearInterval(window._rekorder);
    window._rekorder = setInterval(() => {
        const r = document.getElementById('arbeitsSeiteRahmen'); let fr = '-';
        try { const d = r.contentDocument; if (d && d.location.href !== 'about:blank') fr = d.location.pathname + '?' + (new URLSearchParams(d.location.search).get('ansicht') || ''); } catch (e) {}
        const l = window._arbeitsLive && window._arbeitsLive.lauf;
        const was = l ? (l.was + (l.laeuft ? '' : ' fertig')) : '-';
        const z = window._arbeitsFolgeZiel; const zt = z ? (z.art + ':' + (z.ansicht || '') + ':' + (z.image_id || z.document_id || z.ausgabe_id || '')) : '-';
        let mk = ''; try { const k = r.contentDocument.querySelector('.inkluagent-hin-ziel'); mk = k ? (k.id || k.tagName) : ''; } catch (e) {}
        const e = fr + ' | ' + was + ' | ' + zt + ' | ' + mk;
        if (window._aufnahme[window._aufnahme.length - 1] !== e) window._aufnahme.push(e);
    }, 300); }"""


def fall(pg, name, titel, auftrag, erwartet, verboten, endpunkt):
    print(f"\n=== Fall {name}: {titel}")
    pdf_vorher = {a["id"] for a in (pg.request.get(B + "/api/ausgaben").json().get("ausgaben") or [])}
    r = pg.request.post(B + "/api/upload", multipart={"file": {"name": titel.lower().replace(" ", "-") + ".docx",
                        "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "buffer": word(titel)}})
    pid = (r.json().get("project_id") or r.json().get("id")) if r.ok else None
    check(f"{name}: fiktives Word-Dokument hochgeladen", bool(pid), (r.status, r.text()[:200]))
    neue_ausgaben = set()
    try:
        if not pid:
            return
        for _i in range(60):
            st = pg.request.get(B + f"/api/projects/{pid}/status").json()
            if st.get("status") not in ("extracting", "processing", None):
                break
            time.sleep(3)
        check(f"{name}: zwei Bilder erkannt", (st.get("total_images") or 0) == 2, st)
        pg.goto(B + f"/app?projekt={pid}", wait_until="networkidle")
        pg.wait_for_timeout(3000)
        pg.evaluate(REKORDER)
        pg.evaluate("() => document.querySelectorAll('#inkluagentLog .inkluagent-message').forEach(m => m._alt = true)")
        pg.fill("#inkluagentInput", auftrag)
        pg.click("#inkluagentSendBtn")
        offen = KARTE + "[data-zustand=offen]"
        try:
            pg.wait_for_selector(offen, timeout=180000)
            da = True
        except Exception:  # noqa: BLE001
            da = False
        antwort = pg.evaluate("([...document.querySelectorAll('#inkluagentLog .inkluagent-message.assistant')].pop() || {}).textContent || ''")
        check(f"{name}: das echte Modell bringt eine Auftragskarte", da, antwort[:400])
        if not da:
            return
        pg.wait_for_timeout(2000)
        kid = pg.evaluate(f"() => document.querySelector('{offen}').dataset.kartenId")
        diese = f"{KARTE}[data-karten-id='{kid}']"
        zeilen = pg.evaluate(f"() => [...document.querySelectorAll(\"{diese} .inkluagent-karte-zeilen li\")].map(li => li.textContent)")
        print("  Karte: " + " | ".join(zeilen))
        haken = f"{diese} input[name=per_mail]"
        if pg.locator(haken).count() and pg.is_checked(haken):
            pg.uncheck(haken)
        pg.click(f"{diese} button[type=submit]")
        NEU = ("[...document.querySelectorAll('#inkluagentLog .inkluagent-message.assistant[data-nachricht-id]')]"
               ".filter(m => !m._alt && !m.querySelector('.inkluagent-formular'))")
        t0 = time.time()
        bilder_fahrt = False
        while time.time() - t0 < 600:
            pg.wait_for_timeout(3000)
            if not bilder_fahrt and pg.evaluate("() => (window._aufnahme || []).some(e => e.indexOf('?alttexte | alt_texte | bild:') >= 0)"):
                bilder_fahrt = True
                pg.locator("#arbeitsfenster").screenshot(path=f"/tmp/kette_{name}_alttexte.png")
            if pg.evaluate(f"() => {NEU}.length >= 1"):
                break
        pg.wait_for_timeout(6000)   # die letzte Nachfrage des Live-Fensters und die Fahrt zum Ergebnis
        pg.locator("#arbeitsfenster").screenshot(path=f"/tmp/kette_{name}_ende.png")
        gemeldet = pg.evaluate(f"() => {NEU}.length >= 1")
        aufnahme = pg.evaluate("() => window._aufnahme")
        print("  Aufnahme (Seite im Fenster | Live | Ziel | markiert):")
        for e in aufnahme:
            print("    " + e)
        check(f"{name}: der Agent meldet sich, wenn der Auftrag fertig ist", gemeldet, round(time.time() - t0))
        live_was = {e.split(" | ")[1].replace(" fertig", "") for e in aufnahme} - {"-"}
        ansichten = {e.split(" | ")[0] for e in aufnahme} - {"-"}
        print(f"  Stationen im Live-Fenster: {sorted(live_was)}; Seiten im Fenster: {sorted(ansichten)}")
        for station in erwartet:
            check(f"{name}: Station „{station}“ im Live-Fenster", station in live_was, sorted(live_was))
        for station in verboten:
            check(f"{name}: KEINE Station „{station}“ (nur, was der Agent tut)", station not in live_was, sorted(live_was))
        check(f"{name}: Kamera bei den Bildern in „Alt-Texte“", bilder_fahrt or any("?alttexte | alt_texte | bild:" in e for e in aufnahme),
              aufnahme[:12])
        nach_start = aufnahme[next((i for i, e in enumerate(aufnahme) if "?alttexte" in e), 0):]   # ab der ersten Station
        check(f"{name}: Fenster ging danach zur Station „{endpunkt[0]}“", any(e.startswith(endpunkt[0]) for e in nach_start), nach_start[-5:])
        nachher = {a["id"]: a for a in (pg.request.get(B + "/api/ausgaben").json().get("ausgaben") or [])}
        neue_ausgaben = set(nachher) - pdf_vorher
        print(f"  Neue Ablage-Einträge: {[(i, nachher[i].get('dateiname')) for i in sorted(neue_ausgaben)]}")
        pfad = pg.evaluate("() => { const d = document.getElementById('arbeitsSeiteRahmen').contentDocument; return d ? d.location.pathname : ''; }")
        if endpunkt[1].startswith("."):   # Ergebnis in der Ablage
            markiert = [i for i in neue_ausgaben if any(e.startswith("/ablage?") and e.endswith(f"| ausgabe-{i}") for e in aufnahme)]
            check(f"{name}: am Ende steht „Meine Ablage“ im Fenster, das fertige Ergebnis markiert ({endpunkt[1]})",
                  pfad == "/ablage" and markiert and any(str(nachher[i].get("dateiname", "")).lower().endswith(endpunkt[1]) for i in markiert),
                  (pfad, markiert, [nachher[i].get("dateiname") for i in neue_ausgaben], aufnahme[-4:]))
        else:                             # ohne Datei in der Ablage: die Station des letzten Schritts
            check(f"{name}: ohne Datei in der Ablage steht am Ende die Station des letzten Schritts ({endpunkt[1]}) im Fenster",
                  aufnahme[-1].startswith(endpunkt[1]) and any(e.startswith(endpunkt[1]) and e.endswith("| projectName") for e in nach_start),
                  aufnahme[-4:])
        aussen = pg.evaluate("() => document.activeElement && (document.activeElement.id || document.activeElement.tagName)")
        check(f"{name}: der Fokus blieb draußen", aussen != "arbeitsSeiteRahmen", aussen)
    finally:
        if pid:
            pg.request.delete(B + f"/api/projects/{pid}")
        for i in neue_ausgaben:
            pg.request.delete(B + f"/api/ausgaben/{i}")


with sync_playwright() as p:
    br = p.chromium.launch()
    ctx = br.new_context(locale="de-DE", viewport={"width": 1440, "height": 900})
    pg = ctx.new_page()
    js = []
    pg.on("pageerror", lambda e: js.append(str(e)))
    check("Anmeldung", pg.request.post(B + "/api/login", data={"email": sys.argv[1], "password": sys.argv[2]}).ok)
    vorher = pg.request.get(B + "/api/me").json()["user"]["oberflaeche"]
    try:
        pg.request.put(B + "/api/me/oberflaeche", data={"oberflaeche": "agent"})
        if "A" in FAELLE:
            fall(pg, "A", "Vereinsbrief Kette (fiktiv)",
                 "Mach dieses Word-Dokument bitte komplett fertig, als einen Auftrag: zuerst die Alt-Texte für alle Bilder, danach "
                 "in eine barrierefreie PDF umwandeln. Ich schaue im Arbeitsfenster zu.",
                 erwartet=("alt_texte", "umwandeln"), verboten=("pruefung", "tagging", "uebersetzen", "korrektur"),
                 endpunkt=("/app?dokument", ".pdf"))
        if "B" in FAELLE:
            fall(pg, "B", "Vereinsbrief Gegenfall (fiktiv)",
                 "Bitte nur die Alt-Texte für alle Bilder erzeugen und mir danach die Word-Datei mit den Alt-Texten geben, als einen "
                 "Auftrag. Sonst nichts.",
                 erwartet=("alt_texte", "export"), verboten=("pruefung", "tagging", "uebersetzen", "korrektur", "umwandeln"),
                 endpunkt=("/app?dokument", "/app?dokument"))   # Word ohne ABLAGE_WORD: Link in der Nachricht, keine Ablage
        check("Keine JavaScript-Fehler", not js, js)
    finally:
        pg.request.put(B + "/api/me/oberflaeche", data={"oberflaeche": vorher})
        print("  Ansicht zurueck auf: " + pg.request.get(B + "/api/me").json()["user"]["oberflaeche"])
    br.close()
print(f"ERGEBNIS: {ok} OK, {fehler} FEHLT")
