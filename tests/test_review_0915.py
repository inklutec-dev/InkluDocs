"""Nachpruefung 15.09.2026: Befunde der Selbstpruefung bleiben behoben."""
import os, sys, unittest
HERE = os.path.dirname(os.path.abspath(__file__))
for kandidat in ("/app", os.path.join(os.path.dirname(HERE), "backend")):
    if os.path.isdir(kandidat) and kandidat not in sys.path:
        sys.path.insert(0, kandidat)
import main  # noqa: E402
import pdf_export  # noqa: E402


class Nachpruefung(unittest.TestCase):
    def test_log_alias_vorhanden(self):
        """main.py nutzt `log.*` im Word-PDF/UA-Weg — muss definiert sein (NameError-Fund 15.09.)."""
        self.assertTrue(hasattr(main, "log"))
        self.assertIs(main.log, main.logger)

    def test_creator_steuerzeichen(self):
        self.assertEqual(pdf_export.creator_normieren("A\x00B\x07C\r\nD"), "A B C D")

    def test_xlsx_seite_robust(self):
        unit = {"images": [{"image_path": "", "alt_text": "x", "status": "done", "alt_text_edited": None,
                            "original_alt": None, "langbeschreibung": "", "page_number": "kaputt"}]}
        from openpyxl import load_workbook
        import io
        ws = load_workbook(io.BytesIO(main._build_xlsx_bytes(unit))).active
        self.assertIsNone(ws["B2"].value)


if __name__ == "__main__":
    unittest.main()
